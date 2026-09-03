import os
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pandas as pd

data_path = "P:/3026008.01/Data/"
folder_path = "C:/Users/wenlou/Documents/Python Scripts/output/by_sub"

work_path = "C:/Users/wenlou/Documents/Python Scripts"
os.chdir(work_path)
sub_id_lst = pd.read_pickle(os.path.join(data_path, "exp1_sub_lst.pkl"))

img_names = ["_com_ratio.png", "_com_confi.png", "_AV_raw.png",  \
             "_VA_raw.png", "_VAE.png", "_VAE_com_dis_rec.png", "_VAE_com_dis_rec_op.png","_VAE_com_dis.png"]

image_width = Inches(6)
image_height = Inches(4)
slide_width = Inches(10)
slide_height = Inches(10)
# Create PowerPoint presentation
prs = Presentation()
for sub_id in sub_id_lst:
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    left = top = Inches(0)
    for img_id in list(range(8)):
        if img_id in [2, 5]:
            slide = prs.slides.add_slide(prs.slide_layouts[5])
            left = top = Inches(0)
        elif  img_id in [4, 7] :
            left = Inches(5)
            top = Inches(1.5)        
        image_path = os.path.join(folder_path, 'sub_' +str(sub_id) + img_names[img_id])
        slide.shapes.add_picture(image_path, left, top, width=image_width, height=image_height)
        #left += image_width
        top += image_height



# Save the PowerPoint presentation
prs.save("grouped_images_v2.pptx")
